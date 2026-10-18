"""Filtered Excel export — one sheet per person, leader first."""

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from echo_app import db
from echo_app.services.reports import date_filter_clauses
from echo_app.services.teams import require_role

HEADER_FILL = PatternFill("solid", fgColor="FF360568")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)
ACCENT_FILL = PatternFill("solid", fgColor="FF9AC6C5")
THIN = Side(style="thin", color="FFE4E6EE")
CELL_BORDER = Border(top=THIN, left=THIN, bottom=THIN, right=THIN)

LEADER_HEADERS = ["Date", "Day", "Project", "Leader Work", "Team Discussion", "Team Progress", "Outcome",
                  "Blockers", "Pending Work", "Next Plan", "Chatbot Summary", "Submitted"]
MEMBER_HEADERS = ["Date", "Day", "Project", "Work Done", "Outcome", "Blockers", "Pending Work", "Next Plan",
                  "Chatbot Summary", "Submitted"]
LEADER_WIDTHS = [12, 12, 22, 32, 32, 32, 25, 22, 22, 22, 40, 22]
MEMBER_WIDTHS = [12, 12, 22, 32, 25, 22, 22, 22, 40, 22]


def _sheet_title(name: str, used: set) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", (name or "user")[:28]) or "user"
    title, n = base, 2
    while title.lower() in used:
        title = f"{base[:26]}({n})"
        n += 1
    used.add(title.lower())
    return title


def generate_workbook(team_id, user_ids, filters: dict) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Echo Team Report"

    team = db.fetch_one("SELECT * FROM teams WHERE id = %s", (team_id,))
    sql = """SELECT p.*, tm.role FROM team_members tm
             JOIN profiles p ON p.id = tm.user_id
             WHERE tm.team_id = %s"""
    params = [team_id]
    if user_ids:
        sql += " AND p.id = ANY(%s::uuid[])"
        params.append(list(user_ids))
    sql += " ORDER BY (tm.role = 'leader') DESC, p.username"
    members = db.fetch_all(sql, params)

    used_titles = set()
    for member in members:
        ws = wb.create_sheet(_sheet_title(member["username"], used_titles))
        is_leader = member["role"] == "leader"
        headers = LEADER_HEADERS if is_leader else MEMBER_HEADERS
        total_cols = len(headers)

        # Profile header block
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        title = ws.cell(row=1, column=1, value=f"{member['username']}  ({'Team Leader' if is_leader else 'Member'})")
        title.font = Font(size=16, bold=True, color="FFFFFFFF")
        title.fill = HEADER_FILL
        title.alignment = Alignment(vertical="center", horizontal="left", indent=1)
        ws.row_dimensions[1].height = 28

        labels = [
            ("Team", team["team_name"]),
            ("Username", member["username"]),
            ("Email", member["email"]),
            ("Role", member["role_title"] or ("Team Leader" if is_leader else "Team Member")),
            ("LinkedIn", member["linkedin_url"] or ""),
            ("GitHub", member["github_url"] or ""),
            ("Kaggle", member["kaggle_url"] or ""),
            ("Medium", member["medium_url"] or ""),
        ]
        for i, (label, value) in enumerate(labels):
            row = 2 + i
            c = ws.cell(row=row, column=1, value=label)
            c.font = Font(bold=True)
            c.fill = ACCENT_FILL
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=total_cols)
            ws.cell(row=row, column=2, value=value)

        # Reports table
        table_start = 2 + len(labels) + 1
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=table_start, column=i, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(vertical="center", horizontal="center")
        ws.row_dimensions[table_start].height = 22

        clauses, extra = date_filter_clauses("r", filters)
        reports = db.fetch_all(
            f"""SELECT r.* FROM reports r
                WHERE {' AND '.join(['r.team_id = %s', 'r.user_id = %s'] + clauses)}
                ORDER BY r.report_date ASC, r.submitted_time ASC""",
            [team_id, member["id"]] + extra,
        )

        # One discussion per leader/team/day, so the date alone is the key.
        discussions = {}
        if is_leader:
            d_clauses, d_extra = date_filter_clauses("d", filters)
            for d in db.fetch_all(
                f"""SELECT d.* FROM discussions d
                    WHERE {' AND '.join(['d.team_id = %s', 'd.leader_id = %s'] + d_clauses)}""",
                [team_id, member["id"]] + d_extra,
            ):
                discussions[d["report_date"]] = d["discussion_summary"] or ""

        for idx, r in enumerate(reports):
            day = r["report_date"]
            values = [day.isoformat(), day.strftime("%A"), r["project_name"] or ""]
            if is_leader:
                values += [r["leader_work"] or r["work_done"] or "", discussions.get(day, ""), r["team_progress"] or ""]
            else:
                values.append(r["work_done"] or "")
            submitted = r["submitted_time"]
            values += [
                r["outcome"] or "", r["blockers"] or "", r["pending_work"] or "", r["next_plan"] or "",
                r["chatbot_summary"] or "",
                submitted.astimezone().strftime("%Y-%m-%d %H:%M") if isinstance(submitted, datetime) else "",
            ]
            for col, value in enumerate(values, start=1):
                c = ws.cell(row=table_start + 1 + idx, column=col, value=value)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = CELL_BORDER

        for i, w in enumerate(LEADER_WIDTHS if is_leader else MEMBER_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    if not wb.worksheets:
        wb.create_sheet("Report")
    return wb


def default_filename(team_name: str) -> str:
    safe = re.sub(r"[^a-z0-9]+", "_", team_name or "team", flags=re.I).strip("_") or "team"
    return f"{safe}_report_{datetime.now():%Y%m%d_%H%M%S}.xlsx"


def export_excel(team_id, user_id, filters: dict, path) -> Path:
    """Writes the workbook to `path`. Members can only export their own sheet."""
    role = require_role(team_id, user_id)
    filters = dict(filters or {})
    if role == "member":
        user_ids = [user_id]
    else:
        user_ids = [filters["user_id"]] if filters.get("user_id") else None
    wb = generate_workbook(team_id, user_ids, filters)
    path = Path(path)
    wb.save(path)
    return path
